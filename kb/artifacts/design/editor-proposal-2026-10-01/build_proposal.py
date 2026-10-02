from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE=Path(__file__).resolve().parent
DOC=Document()
sec=DOC.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.72)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=DOC.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(8)
DOC.styles['Normal'].font.size=Pt(11)
DOC.styles['Normal'].paragraph_format.line_spacing=1.13
for name,size in [('Title',31),('Heading 1',24),('Heading 2',14)]:
 DOC.styles[name].font.name='Calibri';DOC.styles[name].font.size=Pt(size)
DOC.styles['Subtitle'].font.size=Pt(12)
DOC.styles['Caption'].font.size=Pt(9)
DOC.styles['Caption'].font.italic=True
DOC.styles['Caption'].font.bold=False
for st in DOC.styles:
 for border in st.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
 for fonts in st.element.xpath('.//w:rFonts'):
  for key in list(fonts.attrib):
   if 'theme' in key.lower(): del fonts.attrib[key]

DOC.core_properties.title='Kabita Live website design proposal'
DOC.core_properties.subject='Editorial preview of the redesigned journal and review priorities'
DOC.core_properties.author='Ahimanikya Satapathy'
DOC.core_properties.keywords='Kabita Live, editorial review, website design proposal'
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('Kabita Live | Editorial preview | ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def p(text,style=None):return DOC.add_paragraph(text,style)
def h(text):return DOC.add_heading(text,2)
def page(title):DOC.add_page_break();DOC.add_heading(title,1)
def fig(name,caption,width=6.83):
 para=DOC.add_paragraph();para.paragraph_format.space_after=Pt(3)
 para.paragraph_format.keep_with_next=True
 run=para.add_run();run.add_picture(str(BASE/'screenshots'/name),width=Inches(width))
 run._r.xpath('.//wp:docPr')[0].set('descr',caption)
 p(caption,'Caption')
def link(text,url):
 para=DOC.add_paragraph();rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rel);run=OxmlElement('w:r');pr=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);t=OxmlElement('w:t');t.text=text;run.append(t);a.append(run);para._p.append(a)

DOC.add_heading('Kabita Live website design proposal',0)
p('A walkthrough for the editors before the website preview','Subtitle')
p('Prepared by Ahimanikya Satapathy for Pradeep Biswal, Editor, and Paresh Kumar Pattnaik, Managing Editor.\n1 October 2026')
p('The redesigned Kabita Live is taking shape as a quiet, connected home for the journal. This proposal introduces the design and reading experience before the website link is shared. Please review whether it represents the journal well, and identify the editorial corrections that should come first.')
fig('home.png','Current local preview of the home page. Website screenshots in this document show work in progress, not a public launch.')
h('What the design is trying to achieve')
p('The home page brings the current edition into view immediately, then offers poems in the journal’s three languages, recent editions and the editors. Warm paper, restrained colour and artwork inspired by Odisha give the publication a recognisable setting without competing with the poems.')
p('The branding cheat sheet already shared remains a quick reference. This document explains how that direction appears across the website, what readers will be able to do, and where editorial judgment is still needed.')

page('Edition navigation and archive')
fig('edition.png','Edition preview showing the date, distinct cover artwork and its accompanying narrative.')
h('From an edition to a poem and its writer')
p('Readers can begin with the current issue, browse earlier editions, or search for a poem or poet. Each edition brings its cover, date and contents together. A poem opens into a dedicated reading page; its byline leads to the writer’s profile and other contributions. This keeps the journal’s people and publication history connected.')
h('An archive that can be explored')
p('The current migration contains 47 edition records, 801 poem records and 429 writer records. These figures describe the captured collection, not a claim that every text, biography or attribution has passed editorial review. Missing sources and conflicting credits remain visible in the review work rather than being silently guessed.')
h('Cover artwork and cultural context')
p('Each edition has a distinct visual idea and a short cover narrative. The series draws on Odisha’s landscapes, materials and everyday life. New artwork is retained separately from historical originals, with source and AI-use records. Please review both the image and its cultural context before approving it for publication.')

page('Poem reading and language versions')
fig('poem.png','Odia poem reader in the current local preview, with language choices and text-size controls.')
h('Readable across three scripts')
p('Odia, Hindi and English use fonts selected for their own scripts. The reader gives the title, poet and verse a clear hierarchy, with generous line spacing and larger-text controls. Source-supported stanza breaks are preserved; ambiguous spacing is held for editorial checking. System, Light and Dark appearance choices help readers use the journal comfortably.')
h('Language versions keep their identity')
p('Where versions are available, readers can move between the original and translations. The original remains the reference. A total of 1,574 AI-assisted translation drafts have been prepared across 787 records; all need independent language review. Another 28 target versions across 14 records are held because originals or attribution details need reconciliation. These drafts must not be treated as approved translations.')
h('Small tools that support reading')
p('Readers can enlarge the text, select passages for saved underlining, and continue to the next poem or more work by the same writer. Sharing keeps the poem’s identity and credit attached. Saved markings are a reading aid, not public comments. Audio experiments remain separate from the promised launch experience.')

page('Writer and editor profiles')
fig('editor.png','Pradeep Biswal’s editor profile illustrates the portrait, biography and selected-books presentation.')
h('Editors and contributors')
p('The home page introduces Pradeep Biswal as Editor and Paresh Kumar Pattnaik as Managing Editor. Dedicated profiles provide space for their literary journeys, selected books and contributions. Writer pages connect biographies and portraits with the poems that appear in the journal, so a byline can become a path to further reading.')
h('Profile accuracy and portrait review')
p('Names, native-script spellings, roles, biography details and book links need confirmation from the editors or reliable sources. Portrait treatments should remain recognisable and respectful. Some images use an AI-assisted artistic treatment; likeness, source credit and permission need review. Missing biographical facts should remain missing until they can be verified.')
h('A quieter ending to each poem')
p('The reading page closes with navigation and the writer connection. Historical contributor notes can be presented in the relevant profile rather than interrupting the verse. Translator and copyright credit remain distinct and must not be removed as if they were ordinary contact details. Uncertain attribution needs editorial resolution.')

page('Mobile reading and Facebook')
h('A practical phone experience')
p('On phones, navigation opens in a large, easy-to-dismiss menu. Edition contents and contribution lists become readable cards; artwork and portraits use the available width. Footer links use labelled icons. The same reading and appearance choices remain available. Please try long Odia lines, language changes and enlarged text when the preview link arrives; those are more useful checks than a screenshot alone.')
h('A consistent visual language')
p('The English Kabita Live masthead and separate Odia signature anchor the identity. Cotton-paper surfaces, dark ink, sea blue and restrained laterite accents carry it across pages. Cormorant Garamond provides the English display character; Source Serif 4, Noto Serif Oriya and Tiro Devanagari Hindi support sustained reading. Artwork creates atmosphere while names and poems remain clear.')
para=DOC.add_paragraph();para.paragraph_format.keep_with_next=True
run=para.add_run();run.add_picture(str(BASE.parent.parent/'social/facebook-2026-09-30/cover-v3.png'),width=Inches(6.83));run._r.xpath('.//wp:docPr')[0].set('descr','Kabita Live Facebook cover with a one-line Odia signature over an illustrated river landscape and an open book.')
p('Current Facebook cover. Original AI-generated illustration, created at Ahimanikya Satapathy’s direction.','Caption')
h('The first social presence')
p('The Facebook Page is live with a branded cover, editor introductions and a pinned welcome post. The website preview links to it from the footer. Adding the website address to Facebook will follow migration. Daily posting is a next-stage plan; a posting calendar and automated publication are not yet in place. Instagram remains deferred.')
link('View the Kabita Live Facebook Page','https://www.facebook.com/profile.php?id=61594946730649')

page('Editorial review before the website link')
p('The design and connected reading pages are available in a local preview. The website is not publicly launched. The next step is an editorial walkthrough, followed by corrections and a separate launch decision. No launch date is proposed in this document.')
h('What to review first')
for text in [
'Identity and voice. Does the design feel appropriate for Kabita Live? Confirm the masthead, signature, editor titles, biographies and portrait likenesses.',
'A complete sample edition. Compare titles, bylines, translator credit, punctuation and stanza breaks with the original publication. Note the edition and poem title for every correction.',
'Language quality. Review Odia shaping and readability, then assess translation drafts with qualified readers. Keep original text, existing credited translations and new drafts clearly distinguished.',
'Reader journeys. Find a poem from an edition, follow its byline, enlarge the text and try the mobile menu. Flag anything confusing or difficult to read.',
'Artwork and credits. Confirm cover narratives, portrait sources and permissions, cultural details and the disclosure of AI assistance.'
]:p(text,'List Bullet')
h('What still needs editorial work')
p('Missing original sources and attribution conflicts need reconciliation. Some biographies and complete review texts remain outstanding. A poem lacks a recorded edition and another lacks an author name. Translation drafts and unresolved stanza structures need language review. These issues should be resolved or explicitly handled before launch; a successful technical build does not settle them.')
p('Reader responses are planned through private correspondence, not public comment threads. Online feedback delivery and consent-based analytics still require their final operational checks before being enabled for readers. Remaining mobile layout studies are proposals until approved and applied.')
h('How to send useful feedback')
p('Please separate factual corrections, reading problems and design preferences. For each item, include the edition or writer name, the exact wording or screen concerned, the proposed correction, and a source where available. A small set of agreed priorities will make the first website review more productive.')
h('What comes next')
p('First, gather comments on this proposal and confirm the editor profiles. Next, share the website preview for a guided review of one complete edition and its reading paths. After the corrections and launch checks, confirm migration and publication separately. A daily social schedule can then be agreed around approved poems, excerpts and credits.')
p('Preview prepared with AI assistance. Screenshots and status reflect the current project records on 1 October 2026. Artwork and poetry remain subject to their respective credits and editorial review.','Caption')
DOC.save(BASE/'Kabita-Live-Design-Proposal-for-Editors.docx')
print(BASE/'Kabita-Live-Design-Proposal-for-Editors.docx')
